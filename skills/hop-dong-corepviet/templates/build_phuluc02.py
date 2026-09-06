# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUT = "/sessions/relaxed-sharp-bell/mnt/outputs/hopdong/PHU-LUC-02-GOI-SETUP-XE.docx"
doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
sec.left_margin = Cm(2.5); sec.right_margin = Cm(2)
st = doc.styles['Normal']
st.font.name = 'Times New Roman'; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.15

def P(text="", bold=False, align='left', size=12, italic=False, sb=0, sa=4, indent=None):
    p = doc.add_paragraph()
    p.alignment = {'left':WD_ALIGN_PARAGRAPH.LEFT,'center':WD_ALIGN_PARAGRAPH.CENTER,
                   'right':WD_ALIGN_PARAGRAPH.RIGHT,'just':WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    p.paragraph_format.space_before = Pt(sb); p.paragraph_format.space_after = Pt(sa)
    if indent: p.paragraph_format.left_indent = Cm(indent)
    r = p.add_run(text); r.bold=bold; r.italic=italic; r.font.size=Pt(size); r.font.name='Times New Roman'
    return p

def ART(t): P(t, bold=True, size=12.5, sb=10, sa=5)
def BUL(t, ind=0.6):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(ind)
    p.paragraph_format.space_after = Pt(3); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("- " + t); r.font.size=Pt(12); r.font.name='Times New Roman'

P("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, align='center')
P("Độc lập - Tự do - Hạnh phúc", bold=True, align='center')
P("---------------o0o---------------", align='center', size=11, sa=12)

P("PHỤ LỤC 02", bold=True, align='center', size=15, sa=2)
P("GHI NHẬN GÓI SETUP XE CHUYÊN NGHIỆP", bold=True, align='center', size=13, sa=2)
P("VÀ ĐIỀU KIỆN SỬ DỤNG BỘ NHẬN DIỆN THƯƠNG HIỆU", bold=True, align='center', size=13, sa=4)
P("(Đính kèm và là bộ phận không tách rời Hợp đồng phân phối hàng hóa số ......../2026/HĐPP-CRV)",
  italic=True, align='center', size=11, sa=12)

P("Hôm nay, ngày ....... tháng ....... năm 2026, tại Thành phố Hồ Chí Minh, hai Bên gồm:", align='just', sa=8)

P("BÊN A (BÊN CUNG CẤP):", bold=True)
P("CÔNG TY TNHH THƯƠNG MẠI - SẢN XUẤT - XUẤT NHẬP KHẨU THUẬN TÍN", bold=True, indent=0.5)
for t in ["Mã số thuế: 0319348507",
          "Địa chỉ: 537/173/36 Đường Tô Ngọc Vân, Phường Thới An, Thành phố Hồ Chí Minh",
          "Người đại diện: Ông HUỲNH NGỌC THUẬN            Chức vụ: Giám đốc"]:
    P(t, indent=0.5, size=11.5, sa=2)

P("BÊN B (BÊN PHÂN PHỐI):", bold=True, sb=8)
for t in ["Ông/Bà (hoặc Hộ kinh doanh/Công ty): ...........................................................................",
          "Số CCCD/Mã số thuế: ...................................  Điện thoại: ..............................................",
          "Địa điểm kinh doanh: ..................................................................................................."]:
    P(t, indent=0.5, size=11.5, sa=2)

P("Hai Bên thống nhất ký Phụ lục 02 với các nội dung sau:", align='just', sb=8, sa=6)

ART("ĐIỀU 1. GHI NHẬN GIAO DỊCH GÓI SETUP XE CHUYÊN NGHIỆP")
P("1.1. Bên B đã đặt mua của Bên A 01 (một) Gói Setup Xe Chuyên Nghiệp mang bộ nhận diện thương hiệu "
  "CỜ RẾP VIỆT (sau đây gọi là \"Xe\"), với thông tin như sau:", align='just')

rows = [("Nội dung","Chi tiết"),
        ("Tên gói","Gói Setup Xe Chuyên Nghiệp - Cờ Rếp Việt"),
        ("Giá trị gói (đã gồm VAT)","....................................... VNĐ"),
        ("Số tiền đặt cọc","....................................... VNĐ"),
        ("Ngày đặt cọc","Ngày ....... tháng ....... năm 2026"),
        ("Số tiền còn lại","....................................... VNĐ"),
        ("Ngày giao xe dự kiến","Ngày ....... tháng ....... năm 2026"),
        ("Địa điểm giao xe","..............................................................."),
        ("Khuyến mãi kèm theo","...............................................................")]
t = doc.add_table(rows=len(rows), cols=2); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
w = [Cm(6.0), Cm(10.0)]
for i,row in enumerate(rows):
    for j,val in enumerate(row):
        c = t.cell(i,j); c.width = w[j]
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
        r = p.add_run(val); r.font.size=Pt(11); r.font.name='Times New Roman'
        if i==0: r.bold=True; p.alignment=WD_ALIGN_PARAGRAPH.CENTER

P("", sa=4)
P("1.2. Danh mục thiết bị, vật dụng trong Gói Setup Xe được liệt kê tại Báo giá/Biên bản bàn giao có "
  "chữ ký của hai Bên, là tài liệu kèm theo Phụ lục này.", align='just')
P("1.3. Thời gian hoàn thiện và giao Xe phụ thuộc tiến độ xưởng sản xuất và lịch khai trương dự kiến của "
  "Bên B. Ngày giao xe dự kiến tại khoản 1.1 là mốc tham chiếu; Bên A thông báo cho Bên B khi có thay đổi.",
  align='just')

ART("ĐIỀU 2. QUYỀN SỞ HỮU XE")
P("2.1. Quyền sở hữu Xe được chuyển từ Bên A sang Bên B kể từ thời điểm Bên B thanh toán đủ 100% giá trị "
  "gói và hai Bên hoàn tất nghiệm thu, bàn giao.", align='just')
P("2.2. Bên B tự chịu trách nhiệm bảo quản, bảo dưỡng, vận hành Xe và mọi rủi ro liên quan đến Xe sau khi "
  "nhận bàn giao.", align='just')
P("2.3. Việc Bên B sở hữu Xe KHÔNG đồng nghĩa với việc Bên B sở hữu hoặc được chuyển giao bất kỳ quyền nào "
  "đối với bộ nhận diện thương hiệu CỜ RẾP VIỆT gắn trên Xe. Quyền sở hữu trí tuệ đối với thương hiệu, logo, "
  "decal, menu, hình ảnh nhận diện vẫn thuộc về Bên A theo Điều 10 của Hợp đồng.", align='just')

ART("ĐIỀU 3. ĐIỀU KIỆN SỬ DỤNG BỘ NHẬN DIỆN THƯƠNG HIỆU TRÊN XE")
P("3.1. Bên B được sử dụng bộ nhận diện thương hiệu CỜ RẾP VIỆT gắn trên Xe trong suốt thời hạn Hợp đồng "
  "phân phối còn hiệu lực, tại địa điểm kinh doanh đã đăng ký.", align='just')
P("3.2. Trong thời gian sử dụng bộ nhận diện, Bên B cam kết:", align='just')
BUL("Chỉ kinh doanh sản phẩm bánh cờ rếp được chế biến từ nguyên vật liệu do Bên A cung cấp theo Hợp đồng phân phối.")
BUL("KHÔNG sử dụng Xe mang nhận diện Cờ Rếp Việt để kinh doanh sản phẩm cùng loại có nguồn gốc từ đơn vị khác, hoặc sản phẩm không thuộc hệ thống Cờ Rếp Việt.")
BUL("KHÔNG tự ý thay đổi, chỉnh sửa, che khuất hoặc thêm bớt nội dung trên bộ nhận diện; không gắn thêm nhãn hiệu, logo của bên thứ ba lên Xe.")
BUL("Giữ gìn Xe và bộ nhận diện sạch sẽ, nguyên vẹn, bảo đảm hình ảnh chuyên nghiệp của hệ thống.")
BUL("Tuân thủ quy định pháp luật về an toàn thực phẩm, trật tự đô thị và các quy định khác tại địa điểm kinh doanh.")
P("3.3. Việc Bên B kinh doanh sản phẩm ngoài hệ thống trên Xe mang nhận diện Cờ Rếp Việt được coi là hành "
  "vi vi phạm nghiêm trọng, làm căn cứ để Bên A đơn phương chấm dứt Hợp đồng phân phối ngay lập tức theo "
  "khoản 14.3 của Hợp đồng.", align='just')

ART("ĐIỀU 4. XỬ LÝ BỘ NHẬN DIỆN KHI HỢP ĐỒNG CHẤM DỨT")
P("4.1. Khi Hợp đồng phân phối chấm dứt vì bất kỳ lý do gì, quyền sử dụng bộ nhận diện thương hiệu CỜ RẾP "
  "VIỆT của Bên B chấm dứt theo.", align='just')
P("4.2. Trong vòng 15 (mười lăm) ngày kể từ ngày chấm dứt, Bên B có nghĩa vụ tự thực hiện và tự chịu chi "
  "phí: tháo gỡ toàn bộ decal, biển hiệu, menu, logo và mọi dấu hiệu nhận diện thương hiệu CỜ RẾP VIỆT "
  "trên Xe và tại điểm bán; ngừng sử dụng tên gọi, hình ảnh Cờ Rếp Việt trên mọi kênh truyền thông, mạng "
  "xã hội và ứng dụng đặt hàng.", align='just')
P("4.3. Bên B vẫn giữ quyền sở hữu phần khung Xe và thiết bị đã thanh toán đủ. Bên A không thu hồi Xe và "
  "không hoàn lại tiền gói setup khi Hợp đồng chấm dứt.", align='just')
P("4.4. Quá thời hạn 15 ngày mà Bên B không thực hiện nghĩa vụ tại khoản 4.2, Bên A có quyền: yêu cầu Bên "
  "B chấm dứt ngay hành vi vi phạm; áp dụng chế tài phạt và bồi thường theo Điều 15 của Hợp đồng; và thực "
  "hiện các biện pháp bảo vệ quyền sở hữu trí tuệ theo quy định pháp luật.", align='just')

ART("ĐIỀU 5. ĐIỀU KHOẢN THI HÀNH")
P("5.1. Phụ lục này là bộ phận không tách rời của Hợp đồng phân phối hàng hóa số ......../2026/HĐPP-CRV. "
  "Các nội dung không quy định tại Phụ lục này được áp dụng theo Hợp đồng.", align='just')
P("5.2. Trường hợp có mâu thuẫn giữa Phụ lục này và Hợp đồng về nội dung liên quan đến Gói Setup Xe và bộ "
  "nhận diện trên Xe, áp dụng theo Phụ lục này.", align='just')
P("5.3. Phụ lục có hiệu lực kể từ ngày hai Bên cùng ký và được lập thành 02 (hai) bản có giá trị pháp lý "
  "như nhau, mỗi Bên giữ 01 (một) bản.", align='just')

doc.add_paragraph()
tb = doc.add_table(rows=2, cols=2); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
for i,txt in enumerate(["ĐẠI DIỆN BÊN A","ĐẠI DIỆN BÊN B"]):
    p = tb.cell(0,i).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); r.bold=True; r.font.size=Pt(12); r.font.name='Times New Roman'
p = tb.cell(1,0).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("(Ký, ghi rõ họ tên và đóng dấu)"); r.italic=True; r.font.size=Pt(10.5); r.font.name='Times New Roman'
for _ in range(3): tb.cell(1,0).add_paragraph()
pn = tb.cell(1,0).add_paragraph(); pn.alignment=WD_ALIGN_PARAGRAPH.CENTER
r = pn.add_run("HUỲNH NGỌC THUẬN"); r.bold=True; r.font.size=Pt(12); r.font.name='Times New Roman'
p = tb.cell(1,1).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("(Ký, ghi rõ họ tên)"); r.italic=True; r.font.size=Pt(10.5); r.font.name='Times New Roman'

doc.save(OUT); print("OK", OUT)
